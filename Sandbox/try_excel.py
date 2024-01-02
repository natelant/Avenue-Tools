# pip install pandas openpyxl
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
import sys

def read_csv_and_create_excel(csv_file, excel_file):
    # Read CSV into pandas DataFrame
    df = pd.read_csv(csv_file)

    # Create a new Excel workbook
    wb = Workbook()

    # Write DataFrame to the first sheet
    ws_data = wb.active
    for row in df.iterrows():
        ws_data.append(row[1].tolist())

    # Create a new sheet for the chart
    ws_chart = wb.create_sheet(title="Chart")

    # Create a bar chart
    chart = BarChart()
    data = Reference(ws_data, min_col=2, min_row=1, max_col=df.shape[1], max_row=df.shape[0])
    categories = Reference(ws_data, min_col=1, min_row=2, max_row=df.shape[0])
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart.title = "Data Chart"

    # Add the chart to the chart sheet
    ws_chart.add_chart(chart, "A1")

    # Save the Excel file
    wb.save(excel_file)

if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Usage: python csv_to_excel_with_chart.py <input_csv_file> <output_excel_file>")
        sys.exit(1)

    input_csv_file = sys.argv[1]
    output_excel_file = sys.argv[2]

    read_csv_and_create_excel(input_csv_file, output_excel_file)

    print(f"Conversion completed. Excel file saved to {output_excel_file}")
